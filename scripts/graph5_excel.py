#!/usr/bin/env python3
"""
Graph 5: AAT vs. log2(L1 SIZE) for three L2 sizes (12 simulations).

L1 = 1, 2, 4, 8 KB, 4-way.  L2 = 16, 32, 64 KB, 8-way.  BLOCKSIZE = 32.
Writes results/graph5.xlsx: miss rates from ./sim, CACTI hit times, AAT as
Excel formulas, and a line chart with one curve per L2 size.

    AAT = HT_L1 + MR_L1 * (HT_L2 + MR_L2 * Miss_Penalty)     (spec Section 9)

Usage (from the project root, after `make`):
    python3 scripts/graph5_excel.py
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "graph5.xlsx")

BLOCKSIZE  = 32
TRACE      = "gcc_trace.txt"
L1_ASSOC, L2_ASSOC = 4, 8
LOG2_L1    = range(10, 14)                       # L1 = 1KB .. 8KB
L2_KB      = [16, 32, 64]                        # one curve per L2 size

# From cacti-spreadsheet.xls, sheet "Miss_Penalty", cell B1.
MISS_PENALTY_NS = 20.1
# From cacti-spreadsheet.xls, sheet "CACTI results", column E, Block Size = 32.
HT_L1_NS = {1: 0.14682, 2: 0.154496, 4: 0.185685, 8: 0.211173}     # L1 4-way
HT_L2_NS = {16: 0.254354, 32: 0.288511, 64: 0.341213}              # L2 8-way


def run(l1_size, l2_size):
    """One simulation; returns (L1 miss rate e, L2 miss rate n)."""
    out = subprocess.run([SIM, str(BLOCKSIZE), str(l1_size), str(L1_ASSOC), str(l2_size), str(L2_ASSOC),
                          "0", "0", TRACE], cwd=TRACES, capture_output=True, text=True, check=True).stdout
    e = float(re.search(r"e\. L1 miss rate:\s+([\d.]+)", out).group(1))
    n = float(re.search(r"n\. L2 miss rate:\s+([\d.]+)", out).group(1))
    return e, n


def main():
    # ---- run the 12 simulations ----
    mr1, mr2 = {}, {}
    for l2 in L2_KB:
        for p in LOG2_L1:
            e, n = run(2 ** p, l2 * 1024)
            mr1[p] = e                    # L1 miss rate does not depend on the L2 below it
            mr2[(p, l2)] = n
            print(f"L1 {2 ** p // 1024} KB  L2 {l2:>2} KB   L1 miss rate = {e:.4f}   L2 miss rate = {n:.4f}")

    # ---- sheet ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Graph 5"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws["A1"] = "Graph 5: AAT vs. log2(L1 SIZE) for three L2 sizes"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = (f"Trace: {TRACE}   BLOCKSIZE = {BLOCKSIZE}   L1 {L1_ASSOC}-way   L2 {L2_ASSOC}-way   "
                "No prefetching   AAT = HT_L1 + MR_L1 x (HT_L2 + MR_L2 x Miss_Penalty)")
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    # inputs: miss penalty + one HT_L2 per L2 size
    inputs = [("Miss_Penalty (ns)", MISS_PENALTY_NS, "sheet Miss_Penalty, cell B1")]
    inputs += [(f"HT_L2 {kb}KB (ns)", HT_L2_NS[kb], f"CACTI results, {kb * 1024} B / 32 B block / 8-way")
               for kb in L2_KB]
    ref = {}
    for i, (name, val, src) in enumerate(inputs):
        r = 3 + i
        ws.cell(row=r, column=1, value=name).font = arial(bold=True)
        c = ws.cell(row=r, column=2, value=val)
        c.font = arial(color="0000FF")                            # blue = hardcoded input
        c.comment = Comment("Source: cacti-spreadsheet.xls, " + src, "graph5_excel.py")
        ref[name] = f"$B${r}"
    MP = ref["Miss_Penalty (ns)"]

    hdr = 9
    groups = [("L1 (4-way)", "D9E1F2", 2),
              ("L2 miss rate (from ./sim)", "DDEBF7", 3),
              ("AAT (ns)", "E2EFDA", 3)]
    headers = (["log2(L1 SIZE)", "L1 SIZE (KB)", "L1 miss rate", "HT_L1 (ns)"]
               + [f"L2 {kb}KB" for kb in L2_KB] + [f"L2 {kb}KB" for kb in L2_KB])
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = arial(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
    col = 3
    for g, color, width in groups:
        ws.merge_cells(start_row=hdr - 1, start_column=col, end_row=hdr - 1, end_column=col + width - 1)
        cell = ws.cell(row=hdr - 1, column=col, value=g)
        cell.font = arial(bold=True)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center")
        col += width

    MR1_COL, HT1_COL, MR2_COL, AAT_COL = 3, 4, 5, 8
    first = hdr + 1
    for i, p in enumerate(LOG2_L1):
        r = first + i
        kb = 2 ** p // 1024
        ws.cell(row=r, column=1, value=p)
        ws.cell(row=r, column=2, value=kb)
        m1 = ws.cell(row=r, column=MR1_COL, value=mr1[p])
        m1.number_format = "0.0000"
        h1 = ws.cell(row=r, column=HT1_COL, value=HT_L1_NS[kb])
        h1.number_format = "0.000000"
        for k, l2 in enumerate(L2_KB):
            m2 = ws.cell(row=r, column=MR2_COL + k, value=mr2[(p, l2)])
            m2.number_format = "0.0000"
            ht2 = ref[f"HT_L2 {l2}KB (ns)"]
            a = ws.cell(row=r, column=AAT_COL + k,
                        value=f"={h1.coordinate}+{m1.coordinate}*({ht2}+{m2.coordinate}*{MP})")
            a.number_format = "0.0000"
        for c in range(1, AAT_COL + len(L2_KB)):
            cell = ws.cell(row=r, column=c)
            cell.font = arial(color="0000FF") if c == HT1_COL else arial()
            cell.border = box
            cell.alignment = Alignment(horizontal="center")
    last = first + len(LOG2_L1) - 1

    aat_rng = f"{ws.cell(row=first, column=AAT_COL).coordinate}:{ws.cell(row=last, column=AAT_COL + 2).coordinate}"
    ws.cell(row=last + 2, column=1, value="Lowest AAT (ns)").font = arial(bold=True)
    low = ws.cell(row=last + 2, column=2, value=f"=MIN({aat_rng})")
    low.number_format = "0.0000"
    low.font = arial(bold=True)
    ws.cell(row=last + 3, column=1,
            value="Blue = hardcoded inputs (CACTI values); black = simulator output or formulas.").font = arial(size=9, color="595959")

    ws.column_dimensions["A"].width = 19
    for c in range(2, AAT_COL + len(L2_KB)):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 13

    # ---- chart: one AAT line per L2 size ----
    chart = ScatterChart()
    chart.title = "AAT vs. log2(L1 SIZE)  (L1 4-way, L2 8-way, BLOCKSIZE = 32)"
    chart.x_axis.title = "log2(L1 SIZE)"
    chart.y_axis.title = "AAT (ns)"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 10, 13
    chart.x_axis.majorUnit = 1
    chart.y_axis.number_format = "0.00"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 11, 20

    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    markers = ["circle", "square", "triangle"]
    colors  = ["2A78D6", "EB6834", "1BAF7A"]
    for k, l2 in enumerate(L2_KB):
        ys = Reference(ws, min_col=AAT_COL + k, min_row=first, max_row=last)
        s = Series(ys, xs, title=f"L2 {l2}KB")
        s.marker.symbol = markers[k]
        s.marker.size = 7
        s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k]
        s.graphicalProperties.line.width = 22000
        s.marker.graphicalProperties.solidFill = colors[k]
        s.marker.graphicalProperties.line.solidFill = colors[k]
        chart.series.append(s)
    ws.add_chart(chart, f"A{last + 5}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    best = min((HT_L1_NS[2 ** p // 1024] + mr1[p] * (HT_L2_NS[l2] + mr2[(p, l2)] * MISS_PENALTY_NS), p, l2)
               for p in LOG2_L1 for l2 in L2_KB)
    print(f"\nLowest AAT: {best[0]:.4f} ns  ->  L1 {2 ** best[1] // 1024} KB 4-way + L2 {best[2]} KB 8-way")
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
