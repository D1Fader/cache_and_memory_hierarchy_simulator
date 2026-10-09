#!/usr/bin/env python3
"""
Graph 1: L1 miss rate vs. log2(L1 SIZE) for 5 associativities (55 simulations).

Runs ./sim for every point, writes the results into results/graph1.xlsx as a
table, and adds an Excel line chart built from that table.

Usage (from the project root, after `make`):
    python3 scripts/graph1_excel.py
Needs: pip install --break-system-packages openpyxl
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "graph1.xlsx")

BLOCKSIZE = 32
TRACE     = "gcc_trace.txt"
LOG2_SIZES = range(10, 21)                       # 1KB .. 1MB
ASSOCS = ["Direct-mapped", "2-way", "4-way", "8-way", "Fully-assoc."]


def assoc_value(label, size):
    return {"Direct-mapped": 1, "2-way": 2, "4-way": 4, "8-way": 8,
            "Fully-assoc.": size // BLOCKSIZE}[label]


def l1_miss_rate(size, assoc):
    """Run one simulation and return measurement e (L1 miss rate)."""
    out = subprocess.run([SIM, str(BLOCKSIZE), str(size), str(assoc), "0", "0", "0", "0", TRACE],
                         cwd=TRACES, capture_output=True, text=True, check=True).stdout
    return float(re.search(r"e\. L1 miss rate:\s+([\d.]+)", out).group(1))


def main():
    # ---- run the 55 simulations ----
    results = {}   # (log2_size, assoc_label) -> miss rate
    for p in LOG2_SIZES:
        size = 2 ** p
        for label in ASSOCS:
            mr = l1_miss_rate(size, assoc_value(label, size))
            results[(p, label)] = mr
            print(f"L1 {size // 1024:>5} KB  {label:<14} miss rate = {mr:.4f}")

    # ---- table ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Graph 1"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")

    ws["A1"] = "Graph 1: L1 miss rate vs. log2(L1 SIZE)"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = f"Trace: {TRACE}   BLOCKSIZE = {BLOCKSIZE}   No L2   No prefetching"
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    headers = ["log2(SIZE)", "L1 SIZE (KB)"] + ASSOCS
    hdr_row = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr_row, column=c, value=h)
        cell.font = arial(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for i, p in enumerate(LOG2_SIZES):
        r = hdr_row + 1 + i
        row = [p, 2 ** p // 1024] + [results[(p, a)] for a in ASSOCS]
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = arial()
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c > 2:
                cell.number_format = "0.0000"
    first, last = hdr_row + 1, hdr_row + len(LOG2_SIZES)

    note = last + 2
    ws.cell(row=note, column=1,
            value="Values = measurement e (L1 miss rate) printed by ./sim; "
                  "fully-associative uses ASSOC = SIZE / BLOCKSIZE.").font = arial(size=9, color="595959")
    for col, w in zip("ABCDEFG", [12, 14, 15, 10, 10, 10, 14]):
        ws.column_dimensions[col].width = w

    # ---- chart: one line per associativity, x = log2(SIZE) ----
    chart = ScatterChart()
    chart.title = "L1 miss rate vs. log2(L1 SIZE)  (gcc, BLOCKSIZE = 32)"
    chart.x_axis.title = "log2(L1 SIZE)"
    chart.y_axis.title = "L1 miss rate"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 10, 20
    chart.x_axis.majorUnit = 1
    chart.y_axis.number_format = "0.00"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 11, 20

    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    markers = ["circle", "square", "triangle", "diamond", "x"]
    colors  = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4"]   # one distinct color per curve
    for k in range(len(ASSOCS)):
        col = 3 + k
        ys = Reference(ws, min_col=col, min_row=hdr_row, max_row=last)
        s = Series(ys, xs, title_from_data=True)
        s.marker.symbol = markers[k]
        s.marker.size = 6
        s.graphicalProperties.line.solidFill = colors[k]
        s.graphicalProperties.line.width = 22000          # ~1.75 pt
        s.marker.graphicalProperties.solidFill = colors[k]
        s.marker.graphicalProperties.line.solidFill = colors[k]
        s.smooth = False
        chart.series.append(s)

    ws.add_chart(chart, f"I{hdr_row}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"\nSaved {OUT}")


if __name__ == "__main__":
    main()
