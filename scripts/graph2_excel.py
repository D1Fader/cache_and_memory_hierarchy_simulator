#!/usr/bin/env python3
"""
Graph 2: AAT vs. log2(L1 SIZE) for DM, 2-, 4-, 8-way (no fully-assoc.).

Runs ./sim for the 44 L1 configurations (same configs as Graph 1, minus FA),
writes the miss rates and CACTI hit times into results/graph2.xlsx, computes
AAT with Excel formulas, and adds an Excel line chart.

    AAT = HT_L1 + MR_L1 * Miss_Penalty          (no L2; spec Section 9)

Usage (from the project root, after `make`):
    python3 scripts/graph2_excel.py
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "graph2.xlsx")

BLOCKSIZE  = 32
TRACE      = "gcc_trace.txt"
LOG2_SIZES = range(10, 21)                       # 1KB .. 1MB
ASSOCS     = [("Direct-mapped", 1), ("2-way", 2), ("4-way", 4), ("8-way", 8)]

# From cacti-spreadsheet.xls, sheet "Miss_Penalty", cell B1.
MISS_PENALTY_NS = 20.1

# From cacti-spreadsheet.xls, sheet "CACTI results", column E "Access Time (ns)",
# rows with Block Size = 32.  {size_KB: {assoc: hit_time_ns}}
# 1KB 8-way is not in the CACTI sheet, so that point has no AAT.
HIT_TIME_NS = {
       1: {1: 0.114797, 2: 0.140329, 4: 0.14682},
       2: {1: 0.12909, 2: 0.161691, 4: 0.154496, 8: 0.180686},
       4: {1: 0.147005, 2: 0.181131, 4: 0.185685, 8: 0.189065},
       8: {1: 0.16383, 2: 0.194195, 4: 0.211173, 8: 0.212911},
      16: {1: 0.198417, 2: 0.223917, 4: 0.233936, 8: 0.254354},
      32: {1: 0.233353, 2: 0.262446, 4: 0.27125, 8: 0.288511},
      64: {1: 0.294627, 2: 0.300727, 4: 0.319481, 8: 0.341213},
     128: {1: 0.3668, 2: 0.374603, 4: 0.38028, 8: 0.401236},
     256: {1: 0.443812, 2: 0.445929, 4: 0.457685, 8: 0.458925},
     512: {1: 0.563451, 2: 0.567744, 4: 0.564418, 8: 0.578177},
    1024: {1: 0.69938, 2: 0.706046, 4: 0.699607, 8: 0.705819},
}


def l1_miss_rate(size, assoc):
    """Run one simulation and return measurement e (L1 miss rate)."""
    out = subprocess.run([SIM, str(BLOCKSIZE), str(size), str(assoc), "0", "0", "0", "0", TRACE],
                         cwd=TRACES, capture_output=True, text=True, check=True).stdout
    return float(re.search(r"e\. L1 miss rate:\s+([\d.]+)", out).group(1))


def main():
    # ---- run the 44 simulations ----
    mr = {}
    for p in LOG2_SIZES:
        size = 2 ** p
        for label, a in ASSOCS:
            mr[(p, a)] = l1_miss_rate(size, a)
            print(f"L1 {size // 1024:>5} KB  {label:<14} miss rate = {mr[(p, a)]:.4f}")

    # ---- sheet ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Graph 2"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws["A1"] = "Graph 2: AAT vs. log2(L1 SIZE)"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = f"Trace: {TRACE}   BLOCKSIZE = {BLOCKSIZE}   No L2   No prefetching   AAT = HT_L1 + MR_L1 x Miss_Penalty"
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    ws["A3"] = "Miss_Penalty (ns)"
    ws["A3"].font = arial(bold=True)
    ws["B3"] = MISS_PENALTY_NS
    ws["B3"].font = arial(color="0000FF")                      # blue = hardcoded input
    ws["B3"].comment = Comment("Source: cacti-spreadsheet.xls, sheet Miss_Penalty, cell B1", "graph2_excel.py")

    # header: log2 | KB | 4 x miss rate | 4 x hit time | 4 x AAT
    hdr = 5
    groups = [("L1 miss rate (from ./sim)", "D9E1F2"),
              ("L1 hit time HT (ns, CACTI)", "FCE4D6"),
              ("AAT (ns) = HT + MR x Miss_Penalty", "E2EFDA")]
    headers = ["log2(SIZE)", "L1 SIZE (KB)"]
    for g, _ in groups:
        headers += [lbl for lbl, _ in ASSOCS]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = arial(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
    for gi, (g, color) in enumerate(groups):          # group labels above
        col = 3 + 4 * gi
        ws.merge_cells(start_row=hdr - 1, start_column=col, end_row=hdr - 1, end_column=col + 3)
        cell = ws.cell(row=hdr - 1, column=col, value=g)
        cell.font = arial(bold=True)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center")

    MR_COL, HT_COL, AAT_COL = 3, 7, 11
    first = hdr + 1
    for i, p in enumerate(LOG2_SIZES):
        r = first + i
        kb = 2 ** p // 1024
        for c, v in ((1, p), (2, kb)):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = arial()
        for k, (_, a) in enumerate(ASSOCS):
            m = ws.cell(row=r, column=MR_COL + k, value=mr[(p, a)])
            m.number_format = "0.0000"
            m.font = arial()
            ht = HIT_TIME_NS[kb].get(a)
            h = ws.cell(row=r, column=HT_COL + k, value=ht if ht is not None else "n/a")
            h.number_format = "0.000000"
            h.font = arial(color="0000FF")                     # blue = hardcoded input
            aat = ws.cell(row=r, column=AAT_COL + k)
            aat.font = arial()
            if ht is not None:
                aat.value = f"={h.coordinate}+{m.coordinate}*$B$3"
                aat.number_format = "0.0000"
            else:                                              # no CACTI entry -> blank -> gap in chart
                h.comment = Comment("No CACTI entry for 1KB 8-way, BLOCKSIZE 32", "graph2_excel.py")
        for c in range(1, AAT_COL + 4):
            cell = ws.cell(row=r, column=c)
            cell.border = box
            cell.alignment = Alignment(horizontal="center")
    last = first + len(LOG2_SIZES) - 1

    # lowest AAT
    aat_rng = f"{ws.cell(row=first, column=AAT_COL).coordinate}:{ws.cell(row=last, column=AAT_COL + 3).coordinate}"
    ws.cell(row=last + 2, column=1, value="Lowest AAT (ns)").font = arial(bold=True)
    low = ws.cell(row=last + 2, column=2, value=f"=MIN({aat_rng})")
    low.number_format = "0.0000"
    low.font = arial(bold=True)
    ws.cell(row=last + 3, column=1,
            value="Blank AAT = no CACTI hit time for that configuration (1KB 8-way).").font = arial(size=9, color="595959")
    ws.cell(row=last + 4, column=1,
            value="Blue = hardcoded inputs (CACTI values); black = simulator output or formulas.").font = arial(size=9, color="595959")

    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 13
    for c in range(3, AAT_COL + 4):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 15

    # ---- chart: one AAT line per associativity ----
    chart = ScatterChart()
    chart.title = "AAT vs. log2(L1 SIZE)  (gcc, BLOCKSIZE = 32, no L2)"
    chart.x_axis.title = "log2(L1 SIZE)"
    chart.y_axis.title = "AAT (ns)"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 10, 20
    chart.x_axis.majorUnit = 1
    chart.y_axis.number_format = "0.0"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 11, 20

    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    markers = ["circle", "square", "triangle", "diamond"]
    colors  = ["2A78D6", "EB6834", "1BAF7A", "EDA100"]
    for k, (label, _) in enumerate(ASSOCS):
        ys = Reference(ws, min_col=AAT_COL + k, min_row=first, max_row=last)
        s = Series(ys, xs, title=label)
        s.marker.symbol = markers[k]
        s.marker.size = 6
        s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k]
        s.graphicalProperties.line.width = 22000
        s.marker.graphicalProperties.solidFill = colors[k]
        s.marker.graphicalProperties.line.solidFill = colors[k]
        chart.series.append(s)
    ws.add_chart(chart, f"A{last + 6}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    # ---- print the answer to the template question ----
    best = min(((HIT_TIME_NS[2 ** p // 1024][a] + mr[(p, a)] * MISS_PENALTY_NS, p, lbl)
                for p in LOG2_SIZES for lbl, a in ASSOCS if a in HIT_TIME_NS[2 ** p // 1024]))
    print(f"\nLowest AAT: {best[0]:.4f} ns  ->  {2 ** best[1] // 1024} KB, {best[2]}")
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
