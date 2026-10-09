#!/usr/bin/env python3
"""
Table #1 (ECE 563): stream buffers on the streaming microbenchmark.

    for (i = 0; i < 1000; i++)  c[i] = a[i] + b[i];     // trace: streams_trace.txt

Part A (the report table): L1 1KB direct-mapped, BLOCKSIZE 16, no L2,
        PREF_N = 0..4 stream buffers, PREF_M = 4.
Part B (question 1): other L1 sizes/associativities, prefetch disabled.
Part C (question 3): BLOCKSIZE 16 vs 32, prefetch disabled.

Writes results/table1.xlsx with all three tables and a bar chart for Part A.

Usage (from the project root, after `make`):
    python3 scripts/table1_excel.py
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "table1.xlsx")
TRACE  = "streams_trace.txt"


def run(bs, size, assoc, n=0, m=0):
    """One L1-only simulation; returns the measurements we need."""
    out = subprocess.run([SIM, str(bs), str(size), str(assoc), "0", "0", str(n), str(m), TRACE],
                         cwd=TRACES, capture_output=True, text=True, check=True).stdout
    g = lambda lbl: float(re.search(lbl + r":\s+([\d.]+)", out).group(1))
    return {"accesses": int(g(r"a\. L1 reads") + g(r"c\. L1 writes")),
            "misses":   int(g(r"b\. L1 read misses") + g(r"d\. L1 write misses")),
            "mr":       g(r"e\. L1 miss rate"),
            "pref":     int(g(r"g\. L1 prefetches")),
            "traffic":  int(g(r"q\. memory traffic"))}


def main():
    # ---- Part A: the report table ----
    part_a = []
    for n in range(0, 5):
        m = 0 if n == 0 else 4
        r = run(16, 1024, 1, n, m)
        part_a.append((n, m, r))
        print(f"PREF_N={n} PREF_M={m}:  miss rate {r['mr']:.4f}  misses {r['misses']}  "
              f"prefetches {r['pref']}  traffic {r['traffic']}")

    # ---- Part B: does L1 size / associativity matter? (prefetch off) ----
    part_b_cfgs = [("1KB direct-mapped", 1024, 1), ("1KB 2-way", 1024, 2), ("1KB 4-way", 1024, 4),
                   ("1KB fully-assoc.", 1024, 64), ("2KB direct-mapped", 2048, 1),
                   ("8KB 4-way", 8192, 4), ("32KB 8-way", 32768, 8)]
    part_b = [(lbl, run(16, size, assoc)) for lbl, size, assoc in part_b_cfgs]

    # ---- Part C: block size 16 vs 32 (prefetch off) ----
    part_c = [(bs, run(bs, 1024, 1)) for bs in (16, 32)]

    # ---- workbook ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    def header(row, labels):
        for c, h in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = arial(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = box

    def body(row, values, fmt=None):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = arial()
            cell.border = box
            cell.alignment = Alignment(horizontal="center")
            if fmt and c in fmt:
                cell.number_format = fmt[c]

    def title(row, text, sub):
        ws.cell(row=row, column=1, value=text).font = arial(bold=True, size=11)
        ws.cell(row=row + 1, column=1, value=sub).font = arial(italic=True, size=9, color="595959")

    ws["A1"] = "Table #1: Stream buffers on the streaming microbenchmark"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = "Trace: streams_trace.txt  (for i<1000: c[i] = a[i] + b[i];  2 loads + 1 store per iteration, 4-byte elements)"
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    # Part A
    title(4, "A. Report table", "L1: 1KB, direct-mapped, BLOCKSIZE 16.  No L2.  PREF_M = 4.")
    header(6, ["PREF_N, PREF_M", "L1 miss rate", "L1 accesses", "L1 misses", "L1 prefetches", "Memory traffic (blocks)"])
    a_first = 7
    for i, (n, m, r) in enumerate(part_a):
        label = "0,0 (pref. disabled)" if n == 0 else f"{n},{m}"
        body(a_first + i, [label, r["mr"], r["accesses"], r["misses"], r["pref"], r["traffic"]], {2: "0.0000"})
    a_last = a_first + len(part_a) - 1

    # Part B
    b_row = a_last + 3
    title(b_row, "B. Question 1: other L1 configurations (prefetch disabled)", "BLOCKSIZE 16.  No L2.")
    header(b_row + 2, ["L1 configuration", "L1 miss rate", "L1 accesses", "L1 misses"])
    for i, (lbl, r) in enumerate(part_b):
        body(b_row + 3 + i, [lbl, r["mr"], r["accesses"], r["misses"]], {2: "0.0000"})
    b_last = b_row + 3 + len(part_b) - 1

    # Part C
    c_row = b_last + 3
    title(c_row, "C. Question 3: block size (prefetch disabled)", "L1: 1KB direct-mapped.  No L2.")
    header(c_row + 2, ["BLOCKSIZE (B)", "L1 miss rate", "L1 accesses", "L1 misses", "Elements per block"])
    for i, (bs, r) in enumerate(part_c):
        body(c_row + 3 + i, [bs, r["mr"], r["accesses"], r["misses"], f"=A{c_row + 3 + i}/4"], {2: "0.0000"})
    c_last = c_row + 3 + len(part_c) - 1
    ws.cell(row=c_last + 1, column=1,
            value="Elements per block = BLOCKSIZE / 4 bytes (uint32_t).").font = arial(size=9, color="595959")

    ws.column_dimensions["A"].width = 22
    for col in "BCDEF":
        ws.column_dimensions[col].width = 15
    ws.row_dimensions[6].height = 30

    # Chart for Part A: miss rate per PREF_N
    chart = BarChart()
    chart.type = "col"
    chart.title = "L1 miss rate vs. number of stream buffers (PREF_M = 4)"
    chart.x_axis.title = "PREF_N, PREF_M"
    chart.y_axis.title = "L1 miss rate"
    chart.y_axis.number_format = "0.000"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend = None
    chart.height, chart.width = 9, 16
    data = Reference(ws, min_col=2, min_row=a_first, max_row=a_last)
    cats = Reference(ws, min_col=1, min_row=a_first, max_row=a_last)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    s = chart.series[0]
    s.graphicalProperties.solidFill = "2A78D6"
    s.graphicalProperties.line.solidFill = "2A78D6"
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showCatName = False
    s.dLbls.showSerName = False
    s.dLbls.showLegendKey = False
    s.dLbls.showPercent = False
    s.dLbls.numFmt = "0.000"
    ws.add_chart(chart, "H4")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    print("\nQuestion 1 (prefetch off):")
    for lbl, r in part_b:
        print(f"  {lbl:<20} miss rate {r['mr']:.4f}")
    print("Question 3 (prefetch off):")
    for bs, r in part_c:
        print(f"  BLOCKSIZE {bs:>2}         miss rate {r['mr']:.4f}  ({r['misses']} misses)")
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
