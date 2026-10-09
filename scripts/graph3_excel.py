#!/usr/bin/env python3
"""
Graph 3: AAT vs. log2(L1 SIZE) with a 16KB 8-way L2 added.

Runs ./sim for L1 = 1, 2, 4, 8 KB x {DM, 2-, 4-, 8-way} with a fixed
16KB 8-way L2 (16 simulations), writes miss rates and CACTI hit times into
results/graph3.xlsx, computes AAT with Excel formulas, and adds a line chart.

    AAT = HT_L1 + MR_L1 * (HT_L2 + MR_L2 * Miss_Penalty)     (spec Section 9)

Usage (from the project root, after `make`):
    python3 scripts/graph3_excel.py
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "graph3.xlsx")

BLOCKSIZE  = 32
TRACE      = "gcc_trace.txt"
LOG2_SIZES = range(10, 14)                       # L1 = 1KB .. 8KB
ASSOCS     = [("Direct-mapped", 1), ("2-way", 2), ("4-way", 4), ("8-way", 8)]
L2_SIZE, L2_ASSOC = 16 * 1024, 8

# From cacti-spreadsheet.xls, sheet "Miss_Penalty", cell B1.
MISS_PENALTY_NS = 20.1
# From cacti-spreadsheet.xls, sheet "CACTI results", column E, row 16384 / 32 / 8.
HT_L2_NS = 0.254354
# From cacti-spreadsheet.xls, sheet "CACTI results", column E, Block Size = 32.
# 1KB 8-way is not in the CACTI sheet, so that point has no AAT.
HIT_TIME_L1_NS = {
    1: {1: 0.114797, 2: 0.140329, 4: 0.14682},
    2: {1: 0.12909, 2: 0.161691, 4: 0.154496, 8: 0.180686},
    4: {1: 0.147005, 2: 0.181131, 4: 0.185685, 8: 0.189065},
    8: {1: 0.16383, 2: 0.194195, 4: 0.211173, 8: 0.212911},
}


def run(size, assoc):
    """One simulation with the 16KB 8-way L2; returns (L1 miss rate e, L2 miss rate n)."""
    out = subprocess.run([SIM, str(BLOCKSIZE), str(size), str(assoc), str(L2_SIZE), str(L2_ASSOC),
                          "0", "0", TRACE], cwd=TRACES, capture_output=True, text=True, check=True).stdout
    e = float(re.search(r"e\. L1 miss rate:\s+([\d.]+)", out).group(1))
    n = float(re.search(r"n\. L2 miss rate:\s+([\d.]+)", out).group(1))
    return e, n


def main():
    # ---- run the 16 simulations ----
    mr1, mr2 = {}, {}
    for p in LOG2_SIZES:
        size = 2 ** p
        for label, a in ASSOCS:
            mr1[(p, a)], mr2[(p, a)] = run(size, a)
            print(f"L1 {size // 1024} KB {label:<14} L1 miss rate = {mr1[(p, a)]:.4f}   L2 miss rate = {mr2[(p, a)]:.4f}")

    # ---- sheet ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Graph 3"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws["A1"] = "Graph 3: AAT vs. log2(L1 SIZE), with a 16KB 8-way L2"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = (f"Trace: {TRACE}   BLOCKSIZE = {BLOCKSIZE}   L2 = 16KB 8-way   No prefetching   "
                "AAT = HT_L1 + MR_L1 x (HT_L2 + MR_L2 x Miss_Penalty)")
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    inputs = [("Miss_Penalty (ns)", MISS_PENALTY_NS, "cacti-spreadsheet.xls, sheet Miss_Penalty, cell B1"),
              ("HT_L2 (ns)", HT_L2_NS, "cacti-spreadsheet.xls, CACTI results, 16384 B / 32 B block / 8-way")]
    for i, (name, val, src) in enumerate(inputs):
        ws.cell(row=3 + i, column=1, value=name).font = arial(bold=True)
        c = ws.cell(row=3 + i, column=2, value=val)
        c.font = arial(color="0000FF")                          # blue = hardcoded input
        c.comment = Comment("Source: " + src, "graph3_excel.py")
    MP, HT2 = "$B$3", "$B$4"

    hdr = 7
    groups = [("L1 miss rate (from ./sim)", "D9E1F2"),
              ("L2 miss rate (from ./sim)", "DDEBF7"),
              ("L1 hit time HT_L1 (ns, CACTI)", "FCE4D6"),
              ("AAT (ns)", "E2EFDA")]
    headers = ["log2(L1 SIZE)", "L1 SIZE (KB)"] + [lbl for _ in groups for lbl, _ in ASSOCS]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = arial(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
    for gi, (g, color) in enumerate(groups):
        col = 3 + 4 * gi
        ws.merge_cells(start_row=hdr - 1, start_column=col, end_row=hdr - 1, end_column=col + 3)
        cell = ws.cell(row=hdr - 1, column=col, value=g)
        cell.font = arial(bold=True)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center")

    MR1_COL, MR2_COL, HT_COL, AAT_COL = 3, 7, 11, 15
    first = hdr + 1
    for i, p in enumerate(LOG2_SIZES):
        r = first + i
        kb = 2 ** p // 1024
        for c, v in ((1, p), (2, kb)):
            ws.cell(row=r, column=c, value=v).font = arial()
        for k, (_, a) in enumerate(ASSOCS):
            m1 = ws.cell(row=r, column=MR1_COL + k, value=mr1[(p, a)])
            m2 = ws.cell(row=r, column=MR2_COL + k, value=mr2[(p, a)])
            for m in (m1, m2):
                m.number_format = "0.0000"
                m.font = arial()
            ht = HIT_TIME_L1_NS[kb].get(a)
            h = ws.cell(row=r, column=HT_COL + k, value=ht if ht is not None else "n/a")
            h.number_format = "0.000000"
            h.font = arial(color="0000FF")
            aat = ws.cell(row=r, column=AAT_COL + k)
            aat.font = arial()
            if ht is not None:
                aat.value = f"={h.coordinate}+{m1.coordinate}*({HT2}+{m2.coordinate}*{MP})"
                aat.number_format = "0.0000"
            else:
                h.comment = Comment("No CACTI entry for 1KB 8-way, BLOCKSIZE 32", "graph3_excel.py")
        for c in range(1, AAT_COL + 4):
            cell = ws.cell(row=r, column=c)
            cell.border = box
            cell.alignment = Alignment(horizontal="center")
    last = first + len(LOG2_SIZES) - 1

    aat_rng = f"{ws.cell(row=first, column=AAT_COL).coordinate}:{ws.cell(row=last, column=AAT_COL + 3).coordinate}"
    ws.cell(row=last + 2, column=1, value="Lowest AAT (ns)").font = arial(bold=True)
    low = ws.cell(row=last + 2, column=2, value=f"=MIN({aat_rng})")
    low.number_format = "0.0000"
    low.font = arial(bold=True)
    ws.cell(row=last + 3, column=1,
            value="Blank AAT = no CACTI hit time for that configuration (1KB 8-way).").font = arial(size=9, color="595959")
    ws.cell(row=last + 4, column=1,
            value="Blue = hardcoded inputs (CACTI values); black = simulator output or formulas.").font = arial(size=9, color="595959")

    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 13
    for c in range(3, AAT_COL + 4):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 15

    # ---- chart ----
    chart = ScatterChart()
    chart.title = "AAT vs. log2(L1 SIZE)  (16KB 8-way L2, BLOCKSIZE = 32)"
    chart.x_axis.title = "log2(L1 SIZE)"
    chart.y_axis.title = "AAT (ns)"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 10, 13
    chart.x_axis.majorUnit = 1
    chart.y_axis.number_format = "0.00"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 11, 20

    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    markers = ["circle", "square", "triangle", "diamond"]
    colors  = ["2A78D6", "EB6834", "1BAF7A", "EDA100"]
    for k, (label, _) in enumerate(ASSOCS):
        ys = Reference(ws, min_col=AAT_COL + k, min_row=first, max_row=last)
        s = Series(ys, xs, title=label)
        s.marker.symbol = markers[k]
        s.marker.size = 7
        s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k]
        s.graphicalProperties.line.width = 22000
        s.marker.graphicalProperties.solidFill = colors[k]
        s.marker.graphicalProperties.line.solidFill = colors[k]
        chart.series.append(s)
    ws.add_chart(chart, f"A{last + 6}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    best = min((HIT_TIME_L1_NS[2 ** p // 1024][a] + mr1[(p, a)] * (HT_L2_NS + mr2[(p, a)] * MISS_PENALTY_NS),
                p, lbl) for p in LOG2_SIZES for lbl, a in ASSOCS if a in HIT_TIME_L1_NS[2 ** p // 1024])
    print(f"\nLowest AAT with L2: {best[0]:.4f} ns  ->  L1 {2 ** best[1] // 1024} KB, {best[2]}")
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
