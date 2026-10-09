#!/usr/bin/env python3
"""
Graph 4: L1 miss rate vs. log2(BLOCKSIZE) for six L1 sizes (24 simulations).

L1 is 4-way; BLOCKSIZE = 16, 32, 64, 128; L1 SIZE = 1, 2, 4, 8, 16, 32 KB.
No L2, no prefetching. Writes results/graph4.xlsx with a table and a line chart.

Usage (from the project root, after `make`):
    python3 scripts/graph4_excel.py
"""
import os, re, subprocess
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIM    = os.path.join(ROOT, "sim")
TRACES = os.path.join(ROOT, "traces")
OUT    = os.path.join(ROOT, "results", "graph4.xlsx")

TRACE       = "gcc_trace.txt"
L1_ASSOC    = 4
LOG2_BLOCKS = range(4, 8)                         # 16, 32, 64, 128 B
SIZES_KB    = [1, 2, 4, 8, 16, 32]                # one curve per L1 size


def l1_miss_rate(blocksize, size):
    """Run one simulation and return measurement e (L1 miss rate)."""
    out = subprocess.run([SIM, str(blocksize), str(size), str(L1_ASSOC), "0", "0", "0", "0", TRACE],
                         cwd=TRACES, capture_output=True, text=True, check=True).stdout
    return float(re.search(r"e\. L1 miss rate:\s+([\d.]+)", out).group(1))


def main():
    # ---- run the 24 simulations ----
    mr = {}
    for kb in SIZES_KB:
        for b in LOG2_BLOCKS:
            mr[(b, kb)] = l1_miss_rate(2 ** b, kb * 1024)
            print(f"L1 {kb:>2} KB  BLOCKSIZE {2 ** b:>3}  miss rate = {mr[(b, kb)]:.4f}")

    # ---- table: rows = block size, columns = L1 size ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Graph 4"
    arial = lambda **kw: Font(name="Arial", **kw)
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    ws["A1"] = "Graph 4: L1 miss rate vs. log2(BLOCKSIZE)"
    ws["A1"].font = arial(bold=True, size=12)
    ws["A2"] = f"Trace: {TRACE}   L1 ASSOC = {L1_ASSOC}   No L2   No prefetching"
    ws["A2"].font = arial(italic=True, size=9, color="595959")

    hdr = 5
    ws.merge_cells(start_row=hdr - 1, start_column=3, end_row=hdr - 1, end_column=2 + len(SIZES_KB))
    g = ws.cell(row=hdr - 1, column=3, value="L1 miss rate (from ./sim), by L1 SIZE")
    g.font = arial(bold=True)
    g.fill = PatternFill("solid", fgColor="D9E1F2")
    g.alignment = Alignment(horizontal="center")

    headers = ["log2(BLOCKSIZE)", "BLOCKSIZE (B)"] + [f"{kb}KB" for kb in SIZES_KB]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr, column=c, value=h)
        cell.font = arial(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = box

    first = hdr + 1
    for i, b in enumerate(LOG2_BLOCKS):
        r = first + i
        row = [b, 2 ** b] + [mr[(b, kb)] for kb in SIZES_KB]
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = arial()
            cell.border = box
            cell.alignment = Alignment(horizontal="center")
            if c > 2:
                cell.number_format = "0.0000"
    last = first + len(LOG2_BLOCKS) - 1

    # best block size for each L1 size (helps answer Q1/Q2)
    best_row = last + 1
    ws.cell(row=best_row, column=1, value="Best BLOCKSIZE").font = arial(bold=True)
    for k in range(len(SIZES_KB)):
        col = 3 + k
        L = ws.cell(row=first, column=col).column_letter
        f = (f"=INDEX($B${first}:$B${last},MATCH(MIN({L}{first}:{L}{last}),{L}{first}:{L}{last},0))")
        cell = ws.cell(row=best_row, column=col, value=f)
        cell.font = arial(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
    ws.cell(row=best_row + 2, column=1,
            value="Best BLOCKSIZE = block size with the lowest miss rate in that column.").font = arial(size=9, color="595959")

    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 15
    for k in range(len(SIZES_KB)):
        ws.column_dimensions[ws.cell(row=1, column=3 + k).column_letter].width = 11

    # ---- chart: one line per L1 size, x = log2(BLOCKSIZE) ----
    chart = ScatterChart()
    chart.title = "L1 miss rate vs. log2(BLOCKSIZE)  (gcc, L1 4-way)"
    chart.x_axis.title = "log2(BLOCKSIZE)"
    chart.y_axis.title = "L1 miss rate"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 4, 7
    chart.x_axis.majorUnit = 1
    chart.y_axis.number_format = "0.00"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 11, 20

    xs = Reference(ws, min_col=1, min_row=first, max_row=last)
    markers = ["circle", "square", "triangle", "diamond", "x", "star"]
    colors  = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]
    for k, kb in enumerate(SIZES_KB):
        ys = Reference(ws, min_col=3 + k, min_row=first, max_row=last)
        s = Series(ys, xs, title=f"{kb}KB")
        s.marker.symbol = markers[k]
        s.marker.size = 7
        s.smooth = False
        s.graphicalProperties.line.solidFill = colors[k]
        s.graphicalProperties.line.width = 22000
        s.marker.graphicalProperties.solidFill = colors[k]
        s.marker.graphicalProperties.line.solidFill = colors[k]
        chart.series.append(s)
    ws.add_chart(chart, f"A{best_row + 4}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    print()
    for kb in SIZES_KB:
        b = min(LOG2_BLOCKS, key=lambda b: mr[(b, kb)])
        print(f"L1 {kb:>2} KB: lowest miss rate {mr[(b, kb)]:.4f} at BLOCKSIZE {2 ** b} B")
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
